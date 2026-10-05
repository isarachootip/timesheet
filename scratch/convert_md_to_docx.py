import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='CBD5E1')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_xml = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", "4")}" w:space="0" w:color="{edge_data.get("color", "CBD5E1")}"/>'
            tcBorders.append(parse_xml(b_xml))
        else:
            b_xml = f'<w:{edge} {nsdecls("w")} w:val="none"/>'
            tcBorders.append(parse_xml(b_xml))
    tcPr.append(tcBorders)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    render_inline_formatting(p, text)
    return p

def render_inline_formatting(paragraph, text, default_font_size=11, default_color=RGBColor(30, 41, 59)):
    # Regex to tokenize bold, italic, code, and text
    # e.g. **bold**, *italic*, `code`
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(default_font_size)
            run.font.color.rgb = default_color
            run.font.name = 'Calibri'
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.size = Pt(default_font_size)
            run.font.color.rgb = default_color
            run.font.name = 'Calibri'
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(default_font_size - 1)
            run.font.color.rgb = RGBColor(194, 65, 12) # orange/brick
        else:
            run = paragraph.add_run(token)
            run.font.size = Pt(default_font_size)
            run.font.color.rgb = default_color
            run.font.name = 'Calibri'

def convert_markdown_to_docx(md_path, docx_path, doc_title="NexTime Documentation"):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    
    # Page setup - Margins (0.8 inch / 2 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    in_code_block = False
    code_block_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        
        # Filter out separator lines (|---|---|)
        data_rows = []
        for r in table_rows:
            # Check if separator row
            cells = [c.strip() for c in r.strip().strip('|').split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            data_rows.append(cells)
        
        if not data_rows:
            table_rows = []
            return

        num_cols = max(len(r) for r in data_rows)
        # Pad shorter rows
        for r in data_rows:
            while len(r) < num_cols:
                r.append('')

        table = doc.add_table(rows=len(data_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for row_idx, row_data in enumerate(data_rows):
            is_header = (row_idx == 0)
            row = table.rows[row_idx]
            for col_idx, cell_text in enumerate(row_data):
                cell = row.cells[col_idx]
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                
                if is_header:
                    set_cell_background(cell, '1E3A8A') # Deep Blue 900
                    set_cell_border(cell, bottom=dict(sz=12, val='single', color='1E40AF'))
                else:
                    bg = 'F8FAFC' if row_idx % 2 == 1 else 'FFFFFF'
                    set_cell_background(cell, bg)
                    set_cell_border(cell, 
                                    bottom=dict(sz=4, val='single', color='E2E8F0'),
                                    top=dict(sz=4, val='single', color='E2E8F0'),
                                    left=dict(sz=4, val='single', color='E2E8F0'),
                                    right=dict(sz=4, val='single', color='E2E8F0'))

                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                
                if is_header:
                    render_inline_formatting(p, cell_text, default_font_size=10, default_color=RGBColor(255, 255, 255))
                    for r in p.runs:
                        r.bold = True
                else:
                    render_inline_formatting(p, cell_text, default_font_size=9.5, default_color=RGBColor(51, 65, 85))

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        table_rows = []

    def flush_code_block():
        nonlocal code_block_lines
        if not code_block_lines:
            return
        
        # Create a shaded single-cell table for the code block
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.rows[0].cells[0]
        cell.width = Inches(6.9)
        set_cell_background(cell, '0F172A') # Slate 900 dark background
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_border(cell, 
                        top=dict(sz=8, val='single', color='334155'),
                        bottom=dict(sz=8, val='single', color='334155'),
                        left=dict(sz=8, val='single', color='38BDF8'), # Cyan accent border
                        right=dict(sz=8, val='single', color='334155'))

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        
        code_text = "".join(code_block_lines).strip('\n')
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(241, 245, 249) # Light text
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        code_block_lines = []

    idx = 0
    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()

        # Handle Code Block delimiter
        if stripped.startswith('```'):
            if in_code_block:
                flush_code_block()
                in_code_block = False
            else:
                if in_table:
                    flush_table()
                    in_table = False
                in_code_block = True
            idx += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            idx += 1
            continue

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(stripped)
            idx += 1
            continue
        else:
            if in_table:
                flush_table()
                in_table = False

        # Handle Blank lines
        if not stripped:
            idx += 1
            continue

        # Handle Headings
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
        elif stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(30, 58, 138) # Blue 900
        elif stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(3, 105, 161) # Sky 700
        elif stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(71, 85, 105) # Slate 600

        # Handle Blockquotes / Callout Alerts
        elif stripped.startswith('> '):
            q_text = stripped[2:].strip()
            # Check alert type
            border_color = '3B82F6'
            bg_color = 'EFF6FF'
            if '[!NOTE]' in q_text:
                q_text = q_text.replace('[!NOTE]', '📌 NOTE: ')
            elif '[!TIP]' in q_text:
                q_text = q_text.replace('[!TIP]', '💡 TIP: ')
                border_color = '10B981'
                bg_color = 'ECFDF5'
            elif '[!WARNING]' in q_text:
                q_text = q_text.replace('[!WARNING]', '⚠️ WARNING: ')
                border_color = 'F59E0B'
                bg_color = 'FFFBEB'
            elif '[!IMPORTANT]' in q_text:
                q_text = q_text.replace('[!IMPORTANT]', '⭐ IMPORTANT: ')
                border_color = '6366F1'
                bg_color = 'EEF2FF'

            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            cell = table.rows[0].cells[0]
            cell.width = Inches(6.9)
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            set_cell_border(cell, left=dict(sz=24, val='single', color=border_color))
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            render_inline_formatting(p, q_text, default_font_size=10.5, default_color=RGBColor(30, 41, 59))
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # Handle Horizontal Rule
        elif stripped in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)

        # Handle Lists
        elif re.match(r'^[\*\-\+]\s+', stripped):
            item_text = re.sub(r'^[\*\-\+]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            render_inline_formatting(p, item_text)
        elif re.match(r'^\d+\.\s+', stripped):
            item_text = re.sub(r'^\d+\.\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            render_inline_formatting(p, item_text)
        
        # Regular Paragraph
        else:
            add_styled_paragraph(doc, stripped, space_after=5, space_before=0)

        idx += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    doc.save(docx_path)
    print(f"[OK] Successfully converted {md_path} -> {docx_path}")

if __name__ == '__main__':
    base_dir = r"c:\atgv\time_sheet"
    
    docs_to_convert = [
        (os.path.join(base_dir, "docs", "SA_SYSTEM_GUIDE.md"), os.path.join(base_dir, "docs", "NexTime_SA_System_Guide.docx")),
        (os.path.join(base_dir, "docs", "TRAINER_MANUAL.md"), os.path.join(base_dir, "docs", "NexTime_Trainer_Manual.docx")),
        (os.path.join(base_dir, "user_manual.md"), os.path.join(base_dir, "NexTime_User_Manual.docx")),
        (os.path.join(base_dir, "docs", "SA_SYSTEM_GUIDE.md"), os.path.join(base_dir, "NexTime_SA_System_Guide.docx")),
        (os.path.join(base_dir, "docs", "TRAINER_MANUAL.md"), os.path.join(base_dir, "NexTime_Trainer_Manual.docx")),
    ]

    for md_file, docx_file in docs_to_convert:
        if os.path.exists(md_file):
            convert_markdown_to_docx(md_file, docx_file)
        else:
            print(f"⚠️ File not found: {md_file}")
